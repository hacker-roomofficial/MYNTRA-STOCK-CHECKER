"""Production web service: auth, upload -> run -> report -> webhook, restart auto-resume."""
import asyncio
import time

import aiohttp
import pytest
from aiohttp import web
from openpyxl import Workbook, load_workbook

from stockwatch.config import Config
from stockwatch.excel_io import write_sample
from stockwatch.models import Item, Status, StyleResult
from stockwatch.server import create_app
from stockwatch.store import Store
from tests.mock_myntra import make_myntra_app, make_webhook_app, style_payload

KEY = "s3cret-key"


async def _serve(app):
    r = web.AppRunner(app, access_log=None)
    await r.setup()
    s = web.TCPSite(r, "127.0.0.1", 0)
    await s.start()
    return r, f"http://127.0.0.1:{s._server.sockets[0].getsockname()[1]}"


def _cfg(tmp_path, api, hook):
    c = Config()
    c.api_base = c.site_base = api
    c.warm_session = False
    c.concurrency, c.rps, c.max_rps = 100, 2000, 5000
    c.webhook_url, c.webhook_key = f"{hook}/webhook/myntra-stock", "secret"
    c.data_dir = str(tmp_path)
    c.out_dir, c.db_path = str(tmp_path / "out"), str(tmp_path / "out" / "s.db")
    c.profile_dir = str(tmp_path / "profile")
    c.api_key = KEY
    c.max_upload_mb = 1
    return c


def _xlsx(path, ids):
    wb = Workbook()
    ws = wb.active
    ws.append(["style_id", "sizes"])
    for i in ids:
        ws.append([i, ""])
    wb.save(path)


async def _wait_done(sess, base, rid, timeout=90):
    end = time.monotonic() + timeout
    while time.monotonic() < end:
        async with sess.get(f"{base}/api/runs/{rid}", headers={"X-Api-Key": KEY}) as r:
            j = await r.json()
        if j["status"] in ("done", "done_with_errors", "failed"):
            return j
        await asyncio.sleep(0.3)
    raise AssertionError("run did not finish")


def test_refuses_to_start_without_api_key(tmp_path):
    c = Config()
    c.api_key = ""
    c.out_dir = str(tmp_path)
    with pytest.raises(SystemExit):
        create_app(c)


def test_auth(tmp_path):
    async def go():
        cfg = _cfg(tmp_path, "http://127.0.0.1:9", "http://127.0.0.1:9")
        runner, base = await _serve(create_app(cfg))
        try:
            async with aiohttp.ClientSession() as s:
                out = {}
                async with s.get(f"{base}/health") as r:
                    out["health"] = r.status
                async with s.get(f"{base}/") as r:
                    out["no_key"] = r.status
                async with s.get(f"{base}/api/runs", headers={"X-Api-Key": "wrong"}) as r:
                    out["bad_key"] = r.status
                async with s.get(f"{base}/api/runs", headers={"X-Api-Key": KEY}) as r:
                    out["good_key"] = r.status
                async with s.get(f"{base}/", headers={"Authorization": aiohttp.encode_basic_auth("me", KEY)}) as r:  # browser login
                    out["basic"], page = r.status, await r.text()
                async with s.get(f"{base}/api/runs/..%2F..%2Fetc/report", headers={"X-Api-Key": KEY}) as r:
                    out["traversal"] = r.status
                return out, page
        finally:
            await runner.cleanup()

    out, page = asyncio.run(go())
    assert out == {"health": 200, "no_key": 401, "bad_key": 401, "good_key": 200, "basic": 200, "traversal": 404}
    assert "Stock Watch" in page


def test_upload_run_report_and_alert(tmp_path):
    ids = list(range(7_000_000, 7_000_600))
    xl = tmp_path / "up.xlsx"
    _xlsx(xl, ids)

    async def go():
        sink: list = []
        r1, api = await _serve(make_myntra_app(flaky=False))
        r2, hook = await _serve(make_webhook_app(sink))
        cfg = _cfg(tmp_path, api, hook)
        r3, base = await _serve(create_app(cfg))
        h = {"X-Api-Key": KEY}
        try:
            async with aiohttp.ClientSession() as s:
                # bad uploads
                async with s.post(f"{base}/api/runs", headers=h, data={"x": "1"}) as r:
                    bad1 = r.status
                fd = aiohttp.FormData()
                fd.add_field("file", b"hello", filename="evil.exe")
                async with s.post(f"{base}/api/runs", headers=h, data=fd) as r:
                    bad2 = r.status
                fd = aiohttp.FormData()
                fd.add_field("file", b"nothing,here\nabc,def\n", filename="x.csv")
                async with s.post(f"{base}/api/runs", headers=h, data=fd) as r:
                    bad3 = r.status
                fd = aiohttp.FormData()
                fd.add_field("file", b"1" * (2 * 1024 * 1024), filename="big.csv")
                async with s.post(f"{base}/api/runs", headers=h, data=fd) as r:
                    bad4 = r.status

                # good upload
                fd = aiohttp.FormData()
                fd.add_field("mode", "fast")
                fd.add_field("file", open(xl, "rb"), filename="up.xlsx")
                async with s.post(f"{base}/api/runs", headers=h, data=fd) as r:
                    assert r.status == 202
                    created = await r.json()
                job = await _wait_done(s, base, created["run_id"])
                async with s.get(f"{base}/api/runs/{created['run_id']}/report", headers=h) as r:
                    rep_status, rep = r.status, await r.read()
                return (bad1, bad2, bad3, bad4), created, job, rep_status, rep, sink
        finally:
            for r in (r3, r2, r1):
                await r.cleanup()

    bads, created, job, rep_status, rep, sink = asyncio.run(go())
    assert bads[0] == 400 and bads[1] == 400 and bads[2] == 400 and bads[3] in (413, 400)
    assert created["styles"] == 600 and job["status"] == "done"
    assert job["summary"]["total_checked"] == 600 and job["summary"]["errors"] == 0
    exp_oos = sum(1 for i in ids if (p := style_payload(i)) and (i % 97 == 0 or i % 10 == 0))
    assert job["summary"]["partial_oos"] + job["summary"]["all_oos"] == exp_oos
    assert rep_status == 200
    (tmp_path / "r.xlsx").write_bytes(rep)
    wb = load_workbook(tmp_path / "r.xlsx", read_only=True)
    assert sum(1 for _ in wb["Out_of_Stock"].iter_rows()) - 1 == exp_oos
    assert sum(len(b["alerts"]) for b in sink) == exp_oos  # WhatsApp (n8n) got exactly the OOS ones


def test_restart_resumes_unfinished_run(tmp_path):
    """Simulate a Railway redeploy mid-run: job left 'running' with 200/500 already saved."""
    ids = list(range(8_000_000, 8_000_500))
    inputs = tmp_path / "inputs"
    inputs.mkdir()
    xl = inputs / "left.xlsx"
    _xlsx(xl, ids)

    cfg0 = _cfg(tmp_path, "http://127.0.0.1:9", "http://127.0.0.1:9")
    st = Store(cfg0.db_path)
    st.job_create("crashed-run", str(xl), "fast")
    st.job_resume("crashed-run")  # status 'running' = the container died here
    st.save_many("crashed-run", [StyleResult(str(i), Status.OK, name="pre") for i in ids[:200]])
    st.close()

    calls: list[int] = []

    async def api(req: web.Request):
        sid = int(req.match_info["sid"])
        calls.append(sid)
        p = style_payload(sid)
        return web.Response(status=404) if p is None else web.json_response(p)

    async def go():
        sink: list = []
        mock = web.Application()
        mock.router.add_get("/gateway/v2/product/{sid}", api)
        r1, apib = await _serve(mock)
        r2, hook = await _serve(make_webhook_app(sink))
        cfg = _cfg(tmp_path, apib, hook)
        r3, base = await _serve(create_app(cfg))  # startup must pick the job up by itself
        try:
            async with aiohttp.ClientSession() as s:
                return await _wait_done(s, base, "crashed-run")
        finally:
            for r in (r3, r2, r1):
                await r.cleanup()

    job = asyncio.run(go())
    assert job["status"] == "done"
    assert job["summary"]["total_checked"] == 500
    assert sorted(calls) == ids[200:], "only the 300 unfinished styles may be fetched again"
