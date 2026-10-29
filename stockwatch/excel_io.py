"""Read style IDs from Excel/CSV, write the result report.

Input columns (header names are case-insensitive, order does not matter):
    style_id   required   (also accepted: styleid, style id, style, id, product id)
    sizes      optional   e.g. "S, M, XL" - only these are watched. Blank = watch all sizes.
A file with no header and just IDs in column A also works.
"""
from __future__ import annotations

import csv
import re
from pathlib import Path
from typing import Iterable, Iterator

from openpyxl import Workbook, load_workbook

from .models import Item

_STYLE_HEADERS = {"styleid", "style_id", "style id", "style", "id", "productid", "product id", "style code"}
_SIZE_HEADERS = {"sizes", "size", "watchsizes", "watch_sizes", "watch sizes"}
_SPLIT = re.compile(r"[,;|/]+")


def _clean_id(v) -> str | None:
    if v is None:
        return None
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    return s if s.isdigit() else None


def _parse_sizes(v) -> tuple[str, ...]:
    if v is None:
        return ()
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return tuple(p.strip() for p in _SPLIT.split(str(v)) if p.strip())


def _rows(path: Path) -> Iterator[tuple]:
    ext = path.suffix.lower()
    if ext in (".csv", ".txt", ".tsv"):
        with open(path, newline="", encoding="utf-8-sig") as f:
            dialect = csv.excel_tab if ext == ".tsv" else csv.excel
            for row in csv.reader(f, dialect):
                yield tuple(row)
    else:
        wb = load_workbook(path, read_only=True, data_only=True)  # streaming: lakhs of rows OK
        try:
            yield from wb.active.iter_rows(values_only=True)
        finally:
            wb.close()


def read_items(path: str | Path) -> tuple[list[Item], int]:
    """Returns (unique items in input order, number of skipped/invalid rows)."""
    path = Path(path)
    rows = _rows(path)
    try:
        first = next(rows)
    except StopIteration:
        return [], 0

    head = [str(c or "").strip().lower() for c in first]
    style_col = next((i for i, h in enumerate(head) if h in _STYLE_HEADERS), None)
    size_col = next((i for i, h in enumerate(head) if h in _SIZE_HEADERS), None)

    pending: Iterable[tuple]
    if style_col is None:  # headerless: first row is data, IDs in column A
        style_col, size_col = 0, (1 if len(first) > 1 else None)
        pending = _chain(first, rows)
    else:
        pending = rows

    # None = "watch every size". A blank sizes cell means all sizes, so if the same style
    # appears twice (once blank, once "M") the union is still ALL sizes - never narrow it.
    merged: dict[str, set[str] | None] = {}
    order: list[str] = []
    skipped = 0
    for row in pending:
        if not row or style_col >= len(row):
            continue
        sid = _clean_id(row[style_col])
        if sid is None:
            if any(c not in (None, "") for c in row):
                skipped += 1
            continue
        sizes = _parse_sizes(row[size_col]) if size_col is not None and size_col < len(row) else ()
        if sid not in merged:
            merged[sid] = set(sizes) if sizes else None
            order.append(sid)
        elif merged[sid] is not None:
            if sizes:
                merged[sid].update(sizes)
            else:
                merged[sid] = None

    items = [Item(sid, () if merged[sid] is None else tuple(sorted(merged[sid]))) for sid in order]
    return items, skipped


def _chain(first, rest):
    yield first
    yield from rest


def write_report(path: str | Path, rows: Iterable[dict]) -> dict[str, int]:
    """Streams rows (dicts from Store.iter_results) into an .xlsx with 3 sheets."""
    wb = Workbook(write_only=True)
    oos = wb.create_sheet("Out_of_Stock")
    err = wb.create_sheet("Errors_NotFound")
    allr = wb.create_sheet("All")
    head = ["style_id", "status", "brand", "name", "oos_sizes", "in_stock_sizes", "link", "error", "checked_at"]
    for ws in (oos, err, allr):
        ws.append(head)

    n = {"oos": 0, "err": 0, "all": 0}
    for r in rows:
        line = [
            r["style_id"], r["status"], r["brand"], r["name"], r["oos_sizes"], r["in_sizes"],
            f"https://www.myntra.com/{r['style_id']}", r["error"], r["checked_at_str"],
        ]
        allr.append(line)
        n["all"] += 1
        if r["status"] in ("PARTIAL_OOS", "ALL_OOS"):
            oos.append(line)
            n["oos"] += 1
        elif r["status"] in ("ERROR", "NOT_FOUND"):
            err.append(line)
            n["err"] += 1
    wb.save(path)
    return n


def write_sample(path: str | Path, count: int = 100, start: int = 10000000) -> None:
    wb = Workbook()
    ws = wb.active
    ws.append(["style_id", "sizes"])
    sizes = ["S, M, L", "M, L, XL", "28, 30, 32", "", "S, M"]
    for i in range(count):
        ws.append([start + i * 7, sizes[i % len(sizes)]])
    wb.save(path)
